"""Workbook field provenance; this catalog contains headers, never reference answers.

Unknown columns remain visible to administrators but are excluded from image-only
extraction until their meaning and form placement have been confirmed.
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook

CATALOG_VERSION = 1
YEAR_PATTERN = re.compile(r"Bastrop County (\d{4}) Clean\.xlsx$")

# Same question, different spreadsheet labels. These IDs are independent of year.
ALIASES = {
    'Sex': 'gender', 'Gender': 'gender',
    'Birth Place': 'birth_place', 'Birthplace': 'birth_place',
    "Father's Birthplace": 'father_birth_place', "Father's Birth Place": 'father_birth_place',
    'Father Birth Place': 'father_birth_place',
    "Mother's Birthplace": 'mother_birth_place', "Mother's Birth Place": 'mother_birth_place',
    'Mother Birth Place': 'mother_birth_place',
    'Relationship': 'relation_to_head', 'Relation to Head': 'relation_to_head',
    'Relation to Head of House': 'relation_to_head',
    'Street': 'street_name', 'Street Address': 'street_name',
    'Number of Dwelling in Order of Visitation': 'dwelling_number',
    'Number of Farm Scheduled': 'number_of_farm_schedule',
    "Father's Mother Tongue": 'father_native_tongue', "Father's Native Tongue": 'father_native_tongue',
    'Tongue of Father': 'father_native_tongue',
    "Mother's Mother Tongue": 'mother_native_tongue', "Mother's Native Tongue": 'mother_native_tongue',
    'Tongue of Mother': 'mother_native_tongue',
}
DERIVED = {'Estimated Birth Year', 'Birth Year', 'Current Address', 'Current Coordinate Point',
           'Coordinates (if applicable)', 'Full Address', 'Mexican', 'Mixed Race'}
# These headers need research-team interpretation or identification of a special
# schedule before they can safely be requested from a population-page image.
UNKNOWN = {'Tribe1', 'Graduation School', 'Recorder District', 'Name',
           'Employment Code', 'Employment Details', 'Employment History',
           'Usual Occupation Code', 'COUNTY', 'STATE', 'Previous Entry',
           'Blood Quantum', 'Clan', 'Ceremonies', 'Floor Construction', 'House Construction',
           'Number of Rooms', 'Lot Number', 'Territorial Citizenship',
           'Read Other Language', 'Speak Other Language', 'Write Other Language',
           'Read English', 'Write English', 'Residence on V-J Day', 'Tribe'}


def canonical_field(header: str) -> str:
    return ALIASES.get(header, re.sub(r'_+', '_', re.sub(r'[^a-z0-9]+', '_', header.casefold())).strip('_'))


def classify_column(header: str | None, year: int) -> str:
    if not header or header.startswith('Unnamed:'):
        return 'unknown'
    if header == 'Line Number':
        return 'identity'
    if header in DERIVED or header == 'Birth Date' or (year == 1910 and header == 'Birth Month') or (year == 1850 and header == 'Industry'):
        # Birth Date is reconstructed from age in these workbooks; it is not the
        # literal birth month/year questions asked in 1900.
        return 'derived'
    if header in UNKNOWN:
        return 'unknown'
    return 'image'


def build_catalog(directory: Path) -> dict:
    workbooks = {}
    for path in sorted(directory.glob('Bastrop County * Clean.xlsx')):
        match = YEAR_PATTERN.fullmatch(path.name)
        if not match:
            continue
        year = int(match.group(1))
        checksum = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = {}
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                raw_headers = next(rows, ())
                headers = [str(v) if v is not None else None for v in raw_headers]
                populated = set()
                row_count = 0
                for row in rows:
                    if any(v is not None for v in row):
                        row_count += 1
                    populated.update(i for i, v in enumerate(row) if v is not None)
                fields = []
                seen = set()
                for index, header in enumerate(headers):
                    if header is None and index not in populated:
                        continue
                    normalized = header.strip() if header else None
                    field_id = canonical_field(normalized) if normalized else f'unnamed_column_{index + 1}'
                    classification = classify_column(normalized, year)
                    issue = None
                    if normalized is None:
                        issue = 'Data under an unnamed header; mapping required.'
                    elif field_id in seen:
                        issue = 'Multiple sheet columns map to this field; resolve before comparison/export.'
                    seen.add(field_id)
                    fields.append({'column_index': index, 'header': header,
                                   'field_id': field_id, 'classification': classification,
                                   'mapping_issue': issue})
                sheets[sheet.title] = {'schedule_type': 'slave' if 'slave' in sheet.title.casefold() else 'population',
                                       'row_count': row_count, 'fields': fields}
        finally:
            workbook.close()
        workbooks[str(year)] = {'file': path.name, 'sha256': checksum, 'sheets': sheets}
    payload = {'version': CATALOG_VERSION, 'workbooks': workbooks}
    payload['catalog_sha256'] = hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()
    return payload
